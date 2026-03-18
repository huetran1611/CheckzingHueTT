import argparse
import csv
import math
import os
import statistics
import subprocess
import sys
import time
from pathlib import Path

import openpyxl


STRATEGIES = [
    ("no_diversification", "test_similarity_no_diversification.py"),
    ("with_diversification", "test_similarity_with_diversification.py"),
]


def _to_float(value):
    if value is None:
        return math.nan
    try:
        return float(value)
    except (TypeError, ValueError):
        return math.nan


def _to_int(value):
    if value is None:
        return -1
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return -1


def read_metrics_from_xlsx(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    try:
        sheet = wb.active
        return {
            "dataset_reported": sheet.cell(row=1, column=1).value,
            "best_fitness": _to_float(sheet.cell(row=1, column=2).value),
            "algorithm_runtime_sec": _to_float(sheet.cell(row=1, column=3).value),
            "best_solution": sheet.cell(row=1, column=4).value,
            "best_t_segment": _to_int(sheet.cell(row=1, column=5).value),
            "segments": _to_int(sheet.cell(row=1, column=6).value),
        }
    finally:
        wb.close()


def run_one_job(
    repo_root,
    out_dir,
    dataset_name,
    run_index,
    strategy_name,
    strategy_script,
    strategy_index,
    args,
):
    iteration_id = run_index * 10 + strategy_index
    result_xlsx = repo_root / (
        f"Random_{args.number_of_cities}_{dataset_name}_{args.segment}_iter-_{iteration_id}_CL2.xlsx"
    )
    if result_xlsx.exists():
        result_xlsx.unlink()

    log_file = out_dir / "logs" / f"{Path(dataset_name).stem}__run{run_index}__{strategy_name}.log"
    env = os.environ.copy()
    env.update(
        {
            "NUMBER_OF_CITIES": str(args.number_of_cities),
            "DATA_SET": dataset_name,
            "ITERATION": str(iteration_id),
            "SEGMENT": str(args.segment),
            "A": str(args.A),
            "L": str(args.L),
            "TIME_LIMIT": str(args.time_limit),
        }
    )

    wall_start = time.perf_counter()
    with open(log_file, "w", encoding="utf-8") as log_stream:
        completed = subprocess.run(
            [sys.executable, str(repo_root / strategy_script)],
            cwd=repo_root,
            env=env,
            stdout=log_stream,
            stderr=subprocess.STDOUT,
            text=True,
            check=False,
        )
    wall_time = time.perf_counter() - wall_start

    row = {
        "dataset": dataset_name,
        "run": run_index,
        "strategy": strategy_name,
        "status": "ok",
        "return_code": completed.returncode,
        "best_fitness": math.nan,
        "algorithm_runtime_sec": math.nan,
        "wall_time_sec": wall_time,
        "segments": -1,
        "best_t_segment": -1,
        "result_xlsx": str(result_xlsx),
        "log_file": str(log_file),
        "error_message": "",
    }

    if completed.returncode != 0:
        row["status"] = "failed"
        row["error_message"] = f"Process exited with code {completed.returncode}"
        return row

    if not result_xlsx.exists():
        row["status"] = "failed"
        row["error_message"] = "Result workbook was not created."
        return row

    metrics = read_metrics_from_xlsx(result_xlsx)
    row.update(
        {
            "best_fitness": metrics["best_fitness"],
            "algorithm_runtime_sec": metrics["algorithm_runtime_sec"],
            "segments": metrics["segments"],
            "best_t_segment": metrics["best_t_segment"],
        }
    )
    return row


def write_csv(path, rows, fieldnames):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def summarize_by_strategy(rows):
    summary_rows = []
    for strategy_name, _ in STRATEGIES:
        valid = [r for r in rows if r["strategy"] == strategy_name and r["status"] == "ok"]
        if not valid:
            summary_rows.append(
                {
                    "strategy": strategy_name,
                    "n_jobs_ok": 0,
                    "avg_best_fitness": math.nan,
                    "std_best_fitness": math.nan,
                    "avg_algorithm_runtime_sec": math.nan,
                    "avg_wall_time_sec": math.nan,
                    "avg_segments": math.nan,
                }
            )
            continue
        best_fitness = [r["best_fitness"] for r in valid]
        algo_runtime = [r["algorithm_runtime_sec"] for r in valid]
        wall_runtime = [r["wall_time_sec"] for r in valid]
        segments = [r["segments"] for r in valid]
        summary_rows.append(
            {
                "strategy": strategy_name,
                "n_jobs_ok": len(valid),
                "avg_best_fitness": statistics.mean(best_fitness),
                "std_best_fitness": statistics.pstdev(best_fitness) if len(best_fitness) > 1 else 0.0,
                "avg_algorithm_runtime_sec": statistics.mean(algo_runtime),
                "avg_wall_time_sec": statistics.mean(wall_runtime),
                "avg_segments": statistics.mean(segments),
            }
        )
    return summary_rows


def pairwise_compare(rows):
    indexed = {}
    for r in rows:
        key = (r["dataset"], r["run"])
        indexed.setdefault(key, {})[r["strategy"]] = r

    pair_rows = []
    with_better_fitness = 0
    no_better_fitness = 0
    tie_fitness = 0

    for key in sorted(indexed.keys()):
        data = indexed[key]
        if "no_diversification" not in data or "with_diversification" not in data:
            continue
        no_div = data["no_diversification"]
        with_div = data["with_diversification"]
        if no_div["status"] != "ok" or with_div["status"] != "ok":
            continue

        delta_fitness = with_div["best_fitness"] - no_div["best_fitness"]
        delta_runtime = with_div["algorithm_runtime_sec"] - no_div["algorithm_runtime_sec"]
        delta_segments = with_div["segments"] - no_div["segments"]

        if delta_fitness < 0:
            fitness_winner = "with_diversification"
            with_better_fitness += 1
        elif delta_fitness > 0:
            fitness_winner = "no_diversification"
            no_better_fitness += 1
        else:
            fitness_winner = "tie"
            tie_fitness += 1

        pair_rows.append(
            {
                "dataset": key[0],
                "run": key[1],
                "fitness_no_div": no_div["best_fitness"],
                "fitness_with_div": with_div["best_fitness"],
                "delta_fitness_with_minus_no": delta_fitness,
                "runtime_no_div_sec": no_div["algorithm_runtime_sec"],
                "runtime_with_div_sec": with_div["algorithm_runtime_sec"],
                "delta_runtime_with_minus_no_sec": delta_runtime,
                "segments_no_div": no_div["segments"],
                "segments_with_div": with_div["segments"],
                "delta_segments_with_minus_no": delta_segments,
                "fitness_winner": fitness_winner,
            }
        )

    overall = {
        "n_pairs_ok": len(pair_rows),
        "with_div_better_fitness_count": with_better_fitness,
        "no_div_better_fitness_count": no_better_fitness,
        "tie_fitness_count": tie_fitness,
        "avg_delta_fitness_with_minus_no": statistics.mean(
            [r["delta_fitness_with_minus_no"] for r in pair_rows]
        )
        if pair_rows
        else math.nan,
        "avg_delta_runtime_with_minus_no_sec": statistics.mean(
            [r["delta_runtime_with_minus_no_sec"] for r in pair_rows]
        )
        if pair_rows
        else math.nan,
        "avg_delta_segments_with_minus_no": statistics.mean(
            [r["delta_segments_with_minus_no"] for r in pair_rows]
        )
        if pair_rows
        else math.nan,
    }
    return pair_rows, overall


def main():
    parser = argparse.ArgumentParser(
        description="Benchmark and compare no-diversification vs diversification strategies."
    )
    parser.add_argument("--data-folder", default="test_data/data_demand_random/50")
    parser.add_argument("--number-of-cities", type=int, default=50)
    parser.add_argument("--runs-per-job", type=int, default=2)
    parser.add_argument("--A", type=int, default=4, help="Drone capacity")
    parser.add_argument("--L", type=int, default=90, help="Drone limit time")
    parser.add_argument("--segment", type=int, default=12)
    parser.add_argument("--time-limit", type=int, default=14000)
    parser.add_argument("--max-files", type=int, default=0, help="0 means run all .dat files")
    parser.add_argument("--output-dir", default="Result/benchmark_similarity")
    args = parser.parse_args()

    repo_root = Path(__file__).resolve().parent
    data_folder = Path(args.data_folder)
    if not data_folder.is_absolute():
        data_folder = (repo_root / data_folder).resolve()
    out_dir = Path(args.output_dir)
    if not out_dir.is_absolute():
        out_dir = (repo_root / out_dir).resolve()
    (out_dir / "logs").mkdir(parents=True, exist_ok=True)

    datasets = sorted(data_folder.glob("*.dat"))
    if args.max_files > 0:
        datasets = datasets[: args.max_files]

    if not datasets:
        raise FileNotFoundError(f"No .dat files found in: {data_folder}")

    all_rows = []
    total_jobs = len(datasets) * args.runs_per_job * len(STRATEGIES)
    completed_jobs = 0

    print("Starting benchmark")
    print(f"Data folder: {data_folder}")
    print(f"Files: {len(datasets)}, runs/job: {args.runs_per_job}, total jobs: {total_jobs}")
    print(f"A={args.A}, L={args.L}, number_of_cities={args.number_of_cities}")

    for dataset_path in datasets:
        dataset_name = dataset_path.name
        for run_index in range(1, args.runs_per_job + 1):
            for strategy_index, (strategy_name, strategy_script) in enumerate(STRATEGIES, start=1):
                completed_jobs += 1
                print(
                    f"[{completed_jobs}/{total_jobs}] "
                    f"dataset={dataset_name} run={run_index} strategy={strategy_name}"
                )
                row = run_one_job(
                    repo_root=repo_root,
                    out_dir=out_dir,
                    dataset_name=dataset_name,
                    run_index=run_index,
                    strategy_name=strategy_name,
                    strategy_script=strategy_script,
                    strategy_index=strategy_index,
                    args=args,
                )
                all_rows.append(row)
                if row["status"] != "ok":
                    print(f"  FAILED: {row['error_message']}")

    raw_csv = out_dir / "raw_results.csv"
    write_csv(
        raw_csv,
        all_rows,
        fieldnames=[
            "dataset",
            "run",
            "strategy",
            "status",
            "return_code",
            "best_fitness",
            "algorithm_runtime_sec",
            "wall_time_sec",
            "segments",
            "best_t_segment",
            "result_xlsx",
            "log_file",
            "error_message",
        ],
    )

    summary_rows = summarize_by_strategy(all_rows)
    summary_csv = out_dir / "summary_by_strategy.csv"
    write_csv(
        summary_csv,
        summary_rows,
        fieldnames=[
            "strategy",
            "n_jobs_ok",
            "avg_best_fitness",
            "std_best_fitness",
            "avg_algorithm_runtime_sec",
            "avg_wall_time_sec",
            "avg_segments",
        ],
    )

    pair_rows, pair_overall = pairwise_compare(all_rows)
    pair_csv = out_dir / "pairwise_comparison.csv"
    write_csv(
        pair_csv,
        pair_rows,
        fieldnames=[
            "dataset",
            "run",
            "fitness_no_div",
            "fitness_with_div",
            "delta_fitness_with_minus_no",
            "runtime_no_div_sec",
            "runtime_with_div_sec",
            "delta_runtime_with_minus_no_sec",
            "segments_no_div",
            "segments_with_div",
            "delta_segments_with_minus_no",
            "fitness_winner",
        ],
    )

    overall_csv = out_dir / "pairwise_overall.csv"
    write_csv(overall_csv, [pair_overall], fieldnames=list(pair_overall.keys()))

    print("\nBenchmark completed.")
    print(f"Raw results: {raw_csv}")
    print(f"Summary by strategy: {summary_csv}")
    print(f"Pairwise comparison: {pair_csv}")
    print(f"Pairwise overall: {overall_csv}")


if __name__ == "__main__":
    main()
