from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def _flatten_metrics(prefix: str, metrics: dict[str, Any]) -> dict[str, Any]:
    return {
        f"{prefix}_solution": json.dumps(metrics.get("solution"), ensure_ascii=False),
        f"{prefix}_multi_visit_trip_count": metrics.get("multi_visit_trip_count"),
        f"{prefix}_avg_customers_per_drone_trip": metrics.get("avg_customers_per_drone_trip"),
        f"{prefix}_avg_demand_per_drone_trip": metrics.get("avg_demand_per_drone_trip"),
        f"{prefix}_drone_trip_times": json.dumps(metrics.get("drone_trip_times", []), ensure_ascii=False),
        f"{prefix}_avg_drone_trip_time": metrics.get("avg_drone_trip_time"),
        f"{prefix}_truck_wait_by_point": json.dumps(metrics.get("truck_wait_by_point", {}), ensure_ascii=False),
        f"{prefix}_drone_wait_by_point": json.dumps(metrics.get("drone_wait_by_point", {}), ensure_ascii=False),
        f"{prefix}_avg_truck_wait_by_point": metrics.get("avg_truck_wait_by_point"),
        f"{prefix}_avg_drone_wait_by_point": metrics.get("avg_drone_wait_by_point"),
    }


def _load_rows(input_dir: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for path in sorted(input_dir.rglob("*.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        row = {
            "job_id": payload.get("job_id"),
            "instance_group": payload.get("instance_group"),
            "instance": payload.get("instance"),
            "instance_name": payload.get("instance_name"),
            "run_index": payload.get("run_index"),
            "seed": payload.get("seed"),
            "drone_capacity": payload.get("drone_capacity"),
            "drone_limit_time": payload.get("drone_limit_time"),
            "status": payload.get("status"),
            "runtime_sec": payload.get("runtime_sec"),
            "terminated_by_time_limit": payload.get("terminated_by_time_limit"),
            "elapsed_seconds": payload.get("elapsed_seconds"),
            "initial_objective": payload.get("initial_objective"),
            "best_fitness": payload.get("best_fitness"),
            "best_multi_visit_fitness": payload.get("best_multi_visit_fitness"),
            "segments_run": payload.get("segments_run"),
            "diversification_rounds": payload.get("diversification_rounds"),
            "weights": json.dumps(payload.get("weights", {}), ensure_ascii=False),
        }
        row.update(_flatten_metrics("best_solution", payload.get("best_solution_metrics", {})))
        row.update(_flatten_metrics("best_multi_visit", payload.get("best_multi_visit_metrics", {})))
        rows.append(row)
    rows.sort(
        key=lambda r: (
            str(r.get("instance_group", "")),
            str(r.get("instance_name", "")),
            float(r.get("drone_capacity", 0) or 0),
            float(r.get("drone_limit_time", 0) or 0),
            int(r.get("run_index", 0) or 0),
        )
    )
    return rows


def _write_excel(rows: list[dict[str, Any]], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "results"

    if not rows:
        ws.append(["message"])
        ws.append(["No result rows found"])
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        return

    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])

    for idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for cell in ws.iter_cols(min_col=idx, max_col=idx, min_row=2, max_row=ws.max_row):
            for item in cell:
                max_len = max(max_len, len(str(item.value)) if item.value is not None else 0)
        ws.column_dimensions[get_column_letter(idx)].width = min(max_len + 2, 60)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Summarize Source_revised2 GitHub artifacts into an Excel file.")
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--output-xlsx", required=True, type=Path)
    parser.add_argument("--output-json", required=True, type=Path)
    args = parser.parse_args()

    rows = _load_rows(args.input_dir)
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    _write_excel(rows, args.output_xlsx)
    print(json.dumps({"rows": len(rows), "xlsx": str(args.output_xlsx), "json": str(args.output_json)}))


if __name__ == "__main__":
    main()
